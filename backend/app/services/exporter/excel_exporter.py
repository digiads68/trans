import asyncio
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .base import BaseExporter, LANG_NAMES
from app.models.schemas import SubtitleEntry


class ExcelExporter(BaseExporter):
    """Export translated subtitles to Excel format."""

    def file_extension(self) -> str:
        return "xlsx"

    async def export(self, entries: list[SubtitleEntry], output_path: str, **options) -> str:
        target_name = LANG_NAMES.get(options.get("target_lang") or "vi", options.get("target_lang") or "vi")

        def _write_excel():
            wb = Workbook()
            ws = wb.active
            ws.title = "Translated Subtitles"

            # Style definitions
            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

            # Headers
            headers = ["#", "Start Time", "End Time", "Original", f"{target_name} Translation", "Status"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            # Column widths
            ws.column_dimensions["A"].width = 6
            ws.column_dimensions["B"].width = 16
            ws.column_dimensions["C"].width = 16
            ws.column_dimensions["D"].width = 50
            ws.column_dimensions["E"].width = 50
            ws.column_dimensions["F"].width = 12

            # Data rows
            for row_idx, entry in enumerate(entries, 2):
                ws.cell(row=row_idx, column=1, value=entry.index).border = thin_border
                ws.cell(row=row_idx, column=2, value=entry.start_time or "").border = thin_border
                ws.cell(row=row_idx, column=3, value=entry.end_time or "").border = thin_border

                orig_cell = ws.cell(row=row_idx, column=4, value=entry.original_text)
                orig_cell.alignment = Alignment(wrap_text=True)
                orig_cell.border = thin_border

                trans_cell = ws.cell(
                    row=row_idx, column=5,
                    value=entry.translated_text or "",
                )
                trans_cell.alignment = Alignment(wrap_text=True)
                trans_cell.border = thin_border
                ws.cell(row=row_idx, column=6, value=entry.status or "").border = thin_border

            wb.save(output_path)

        await asyncio.to_thread(_write_excel)
        return output_path
